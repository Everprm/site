from flask import Flask, request, jsonify, render_template, session, redirect, url_for, send_file
import psycopg2
from psycopg2.extras import RealDictCursor
import bcrypt
from datetime import datetime, timedelta
import os
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from functools import wraps
from dotenv import load_dotenv

# Загружаем переменные из .env
load_dotenv()

app = Flask(__name__)

# ============================================================
# БЕЗОПАСНОСТЬ: секретный ключ из переменной окружения
# ============================================================
SECRET_KEY = os.environ.get('SECRET_KEY')
if not SECRET_KEY:
    raise ValueError(
        "❌ Переменная окружения SECRET_KEY не задана!\n"
        "   Добавьте её в файл .env (локально) или в переменные окружения App Platform (на хостинге).\n"
        "   Сгенерировать можно командой: python -c \"import secrets; print(secrets.token_hex(32))\""
    )
app.secret_key = SECRET_KEY
app.permanent_session_lifetime = timedelta(hours=8)

# ============================================================
# ПОДКЛЮЧЕНИЕ К БД
# ============================================================
DATABASE_URL = os.environ.get('DATABASE_URL')
if not DATABASE_URL:
    raise ValueError(
        "❌ Переменная окружения DATABASE_URL не задана!\n"
        "   Добавьте её в файл .env (локально) или в переменные окружения App Platform (на хостинге)."
    )


# ============================================================
# ПЕРМСКОЕ ВРЕМЯ (UTC+5)
# ============================================================

def get_perm_time():
    """Возвращает текущее пермское время (UTC+5)"""
    return datetime.utcnow() + timedelta(hours=5)


# ============================================================
# БАЗА ДАННЫХ
# ============================================================

def get_db():
    """Подключение к PostgreSQL"""
    conn = psycopg2.connect(DATABASE_URL, cursor_factory=RealDictCursor)
    return conn


# ============================================================
# ДЕКОРАТОРЫ
# ============================================================

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login_page'))
        return f(*args, **kwargs)
    return decorated_function


def role_required(allowed_roles):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if 'role' not in session or session['role'] not in allowed_roles:
                return jsonify({'error': 'Доступ запрещен. Недостаточно прав.'}), 403
            return f(*args, **kwargs)
        return decorated_function
    return decorator


# ============================================================
# ПРАВА ПОЛЬЗОВАТЕЛЕЙ
# ============================================================

CAN_SHIP_USERS = ['Павел', 'Валерий']
CAN_RESERVE_USERS = ['Павел', 'Евгений', 'Виталий']


def can_ship(username):
    return username in CAN_SHIP_USERS


def can_reserve(username):
    return username in CAN_RESERVE_USERS


# ============================================================
# СТРАНИЦЫ
# ============================================================

@app.route('/login')
def login_page():
    if 'user_id' in session:
        return redirect(url_for('index'))
    return render_template('login.html')


@app.route('/')
@login_required
def index():
    return render_template('index.html', username=session.get('username'), role=session.get('role'))


@app.route('/history')
@login_required
def history():
    if session.get('role') != 'admin':
        return render_template('access_denied.html'), 403
    
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT 
            sm.id,
            p.name as product_name,
            sm.quantity,
            sm.user,
            sm.comment,
            to_char(sm.created_at + INTERVAL '5 hours', 'YYYY-MM-DD HH24:MI:SS') as created_at
        FROM stock_moves sm
        JOIN products p ON sm.product_id = p.id
        ORDER BY sm.created_at DESC
        LIMIT 200
    ''')
    logs = cur.fetchall()
    cur.close()
    conn.close()
    return render_template('history.html', logs=logs, username=session.get('username'))


@app.route('/reserves')
@login_required
def reserves_page():
    return render_template('reserves.html', username=session.get('username'))


@app.route('/admin/products')
@login_required
def admin_products():
    if session.get('role') != 'admin':
        return render_template('access_denied.html'), 403
    return render_template('admin_products.html', username=session.get('username'))


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login_page'))


# ============================================================
# API - АВТОРИЗАЦИЯ
# ============================================================

@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.get_json()
    username = data.get('username')
    password = data.get('password')
    
    if not username or not password:
        return jsonify({'error': 'Введите логин и пароль'}), 400
    
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, name, password_hash, role FROM users WHERE name = %s", (username,))
    user = cur.fetchone()
    cur.close()
    conn.close()
    
    if not user:
        return jsonify({'error': 'Пользователь не найден'}), 401
    
    try:
        if bcrypt.checkpw(password.encode('utf-8'), user['password_hash'].encode('utf-8')):
            session.permanent = True
            session['user_id'] = user['id']
            session['username'] = user['name']
            session['role'] = user['role']
            return jsonify({'status': 'OK', 'user': user['name'], 'role': user['role']})
        else:
            return jsonify({'error': 'Неверный пароль'}), 401
    except Exception as e:
        return jsonify({'error': f'Ошибка проверки пароля: {str(e)}'}), 500


@app.route('/api/current_user')
@login_required
def current_user():
    username = session.get('username')
    return jsonify({
        'username': username,
        'role': session.get('role'),
        'can_ship': can_ship(username),
        'can_reserve': can_reserve(username)
    })


@app.route('/api/users')
@login_required
def get_users():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT name FROM users ORDER BY name")
    users = cur.fetchall()
    cur.close()
    conn.close()
    return jsonify([u['name'] for u in users])


# ============================================================
# API - ОСТАТКИ
# ============================================================

@app.route('/api/balances')
@login_required
def balances():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT 
            p.id, 
            p.name, 
            p.unit, 
            COALESCE(SUM(sm.quantity), 0) as balance,
            COALESCE((
                SELECT SUM(r.quantity) 
                FROM reserves r 
                WHERE r.product_id = p.id AND r.is_active = 1
            ), 0) as reserved
        FROM products p
        LEFT JOIN stock_moves sm ON p.id = sm.product_id
        GROUP BY p.id, p.name, p.unit, p.sort_order
        ORDER BY p.sort_order, p.id
    ''')
    data = cur.fetchall()
    cur.close()
    conn.close()
    
    result = []
    for row in data:
        item = dict(row)
        item['available'] = item['balance'] - item['reserved']
        result.append(item)
    
    return jsonify(result)


@app.route('/api/ship', methods=['POST'])
@login_required
def ship():
    username = session.get('username')
    
    if not can_ship(username):
        return jsonify({'error': f'❌ У пользователя {username} нет прав на отгрузку'}), 403

    data = request.get_json()
    prod_id = data.get('product_id')
    qty = data.get('quantity')
    comment = data.get('comment', 'Отгрузка')

    if not prod_id or not qty or qty <= 0:
        return jsonify({'error': 'Некорректные данные'}), 400

    conn = get_db()
    cur = conn.cursor()
    
    cur.execute('''
        SELECT 
            COALESCE((SELECT SUM(quantity) FROM stock_moves WHERE product_id = %s), 0) as balance,
            COALESCE((SELECT SUM(quantity) FROM reserves WHERE product_id = %s AND is_active = 1), 0) as reserved
    ''', (prod_id, prod_id))
    check = cur.fetchone()
    
    balance = check['balance']
    reserved = check['reserved']
    available = balance - reserved
    
    if qty > available:
        cur.close()
        conn.close()
        if reserved > 0:
            return jsonify({
                'error': f'❌ Нельзя отгрузить {qty} шт! Доступно: {available} (остаток: {balance}, в резерве: {reserved})'
            }), 400
        else:
            return jsonify({
                'error': f'❌ Недостаточно товара! Остаток: {balance}'
            }), 400

    cur.execute(
        "INSERT INTO stock_moves (product_id, quantity, \"user\", comment) VALUES (%s, %s, %s, %s)",
        (prod_id, -abs(qty), username, comment)
    )
    conn.commit()
    cur.close()
    conn.close()
    return jsonify({'status': 'OK', 'message': f'Отгружено {qty} единиц'})


@app.route('/api/receive', methods=['POST'])
@login_required
@role_required(['manager', 'admin'])
def receive():
    data = request.get_json()
    prod_id = data.get('product_id')
    qty = data.get('quantity')
    user = session.get('username', 'Неизвестный')
    comment = data.get('comment', 'Пополнение')

    if not prod_id or not qty or qty <= 0:
        return jsonify({'error': 'Некорректные данные'}), 400

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO stock_moves (product_id, quantity, \"user\", comment) VALUES (%s, %s, %s, %s)",
        (prod_id, abs(qty), user, comment)
    )
    conn.commit()
    cur.close()
    conn.close()
    return jsonify({'status': 'OK', 'message': f'Добавлено {qty} единиц'})


# ============================================================
# API - УПРАВЛЕНИЕ ТОВАРАМИ (только для админа)
# ============================================================

@app.route('/api/admin/products', methods=['GET'])
@login_required
@role_required(['admin'])
def admin_get_products():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT 
            p.id, 
            p.name, 
            p.unit, 
            p.sort_order,
            COALESCE(SUM(sm.quantity), 0) as balance,
            COALESCE((
                SELECT SUM(r.quantity) 
                FROM reserves r 
                WHERE r.product_id = p.id AND r.is_active = 1
            ), 0) as reserved,
            (SELECT COUNT(*) FROM stock_moves WHERE product_id = p.id) as moves_count
        FROM products p
        LEFT JOIN stock_moves sm ON p.id = sm.product_id
        GROUP BY p.id, p.name, p.unit, p.sort_order
        ORDER BY p.sort_order, p.id
    ''')
    data = cur.fetchall()
    cur.close()
    conn.close()
    
    result = []
    for row in data:
        item = dict(row)
        item['available'] = item['balance'] - item['reserved']
        result.append(item)
    
    return jsonify(result)


@app.route('/api/admin/products', methods=['POST'])
@login_required
@role_required(['admin'])
def admin_add_product():
    data = request.get_json()
    name = (data.get('name') or '').strip()
    unit = (data.get('unit') or 'шт').strip() or 'шт'
    initial_qty = data.get('initial_qty', 0)
    user = session.get('username', 'admin')
    
    if not name:
        return jsonify({'error': 'Укажите наименование товара'}), 400
    
    try:
        initial_qty = int(initial_qty)
        if initial_qty < 0:
            initial_qty = 0
    except (ValueError, TypeError):
        initial_qty = 0
    
    conn = get_db()
    cur = conn.cursor()
    
    cur.execute("SELECT id FROM products WHERE name = %s", (name,))
    existing = cur.fetchone()
    if existing:
        cur.close()
        conn.close()
        return jsonify({'error': f'Товар с именем "{name}" уже существует'}), 400
    
    # Новый товар получает sort_order = максимальный + 1, чтобы встать в конец
    cur.execute("SELECT COALESCE(MAX(sort_order), 0) + 1 as new_order FROM products")
    new_sort_order = cur.fetchone()['new_order']
    
    cur.execute(
        "INSERT INTO products (name, unit, sort_order) VALUES (%s, %s, %s) RETURNING id",
        (name, unit, new_sort_order)
    )
    new_id = cur.fetchone()['id']
    
    if initial_qty > 0:
        cur.execute(
            "INSERT INTO stock_moves (product_id, quantity, \"user\", comment) VALUES (%s, %s, %s, %s)",
            (new_id, initial_qty, user, 'Начальный остаток')
        )
    
    conn.commit()
    cur.close()
    conn.close()
    
    print(f"✅ Добавлен товар: ID={new_id}, '{name}', {initial_qty} {unit}, sort_order={new_sort_order}")
    
    return jsonify({
        'status': 'OK',
        'message': f'Товар "{name}" добавлен (остаток: {initial_qty} {unit})',
        'product_id': new_id
    })


@app.route('/api/admin/products/<int:product_id>', methods=['PUT'])
@login_required
@role_required(['admin'])
def admin_update_product(product_id):
    data = request.get_json()
    name = (data.get('name') or '').strip()
    unit = (data.get('unit') or 'шт').strip() or 'шт'
    sort_order = data.get('sort_order')  # опционально
    
    if not name:
        return jsonify({'error': 'Укажите наименование товара'}), 400
    
    conn = get_db()
    cur = conn.cursor()
    
    cur.execute("SELECT id FROM products WHERE id = %s", (product_id,))
    product = cur.fetchone()
    if not product:
        cur.close()
        conn.close()
        return jsonify({'error': 'Товар не найден'}), 404
    
    cur.execute("SELECT id FROM products WHERE name = %s AND id != %s", (name, product_id))
    duplicate = cur.fetchone()
    if duplicate:
        cur.close()
        conn.close()
        return jsonify({'error': f'Товар с именем "{name}" уже существует'}), 400
    
    # Если передан sort_order — обновляем, иначе оставляем как было
    if sort_order is not None:
        try:
            sort_order = int(sort_order)
            cur.execute(
                "UPDATE products SET name = %s, unit = %s, sort_order = %s WHERE id = %s",
                (name, unit, sort_order, product_id)
            )
        except (ValueError, TypeError):
            cur.execute(
                "UPDATE products SET name = %s, unit = %s WHERE id = %s",
                (name, unit, product_id)
            )
    else:
        cur.execute(
            "UPDATE products SET name = %s, unit = %s WHERE id = %s",
            (name, unit, product_id)
        )
    
    conn.commit()
    cur.close()
    conn.close()
    
    return jsonify({'status': 'OK', 'message': f'Товар обновлён'})


@app.route('/api/admin/products/<int:product_id>', methods=['DELETE'])
@login_required
@role_required(['admin'])
def admin_delete_product(product_id):
    conn = get_db()
    cur = conn.cursor()
    
    cur.execute("SELECT id, name FROM products WHERE id = %s", (product_id,))
    product = cur.fetchone()
    if not product:
        cur.close()
        conn.close()
        return jsonify({'error': 'Товар не найден'}), 404
    
    cur.execute("SELECT COUNT(*) as cnt FROM stock_moves WHERE product_id = %s", (product_id,))
    moves = cur.fetchone()
    
    cur.execute("SELECT COUNT(*) as cnt FROM reserves WHERE product_id = %s AND is_active = 1", (product_id,))
    reserves = cur.fetchone()
    
    if moves['cnt'] > 1 or reserves['cnt'] > 0:
        cur.close()
        conn.close()
        reasons = []
        if moves['cnt'] > 1:
            reasons.append(f'движений: {moves["cnt"]}')
        if reserves['cnt'] > 0:
            reasons.append(f'активных резервов: {reserves["cnt"]}')
        return jsonify({
            'error': f'Нельзя удалить товар "{product["name"]}" — есть история: {", ".join(reasons)}.'
        }), 400
    
    cur.execute("DELETE FROM stock_moves WHERE product_id = %s", (product_id,))
    cur.execute("DELETE FROM reserves WHERE product_id = %s", (product_id,))
    cur.execute("DELETE FROM products WHERE id = %s", (product_id,))
    conn.commit()
    cur.close()
    conn.close()
    
    return jsonify({'status': 'OK', 'message': f'Товар "{product["name"]}" удалён'})


# ============================================================
# API - РЕЗЕРВЫ
# ============================================================

@app.route('/api/reserve', methods=['POST'])
@login_required
def reserve_product():
    username = session.get('username')
    
    if not can_reserve(username):
        return jsonify({'error': f'❌ У пользователя {username} нет прав на резервирование'}), 403

    data = request.get_json()
    prod_id = data.get('product_id')
    qty = data.get('quantity')
    comment = (data.get('comment') or 'Без пометки').strip() or 'Без пометки'

    if not prod_id or not qty or qty <= 0:
        return jsonify({'error': 'Некорректные данные'}), 400

    conn = get_db()
    cur = conn.cursor()
    
    cur.execute('''
        SELECT 
            COALESCE((SELECT SUM(quantity) FROM stock_moves WHERE product_id = %s), 0) as balance,
            COALESCE((SELECT SUM(quantity) FROM reserves WHERE product_id = %s AND is_active = 1), 0) as reserved
    ''', (prod_id, prod_id))
    check = cur.fetchone()
    
    available = check['balance'] - check['reserved']
    
    if available < qty:
        cur.close()
        conn.close()
        return jsonify({
            'error': f'Недостаточно товара для резерва! Доступно: {available}'
        }), 400

    cur.execute(
        "INSERT INTO reserves (product_id, quantity, \"user\", comment, is_active) VALUES (%s, %s, %s, %s, 1) RETURNING id",
        (prod_id, qty, username, comment)
    )
    new_id = cur.fetchone()['id']
    conn.commit()
    cur.close()
    conn.close()
    
    return jsonify({
        'status': 'OK', 
        'message': f'Зарезервировано {qty} единиц. Пометка: {comment}',
        'reserve_id': new_id
    })


@app.route('/api/unreserve', methods=['POST'])
@login_required
def unreserve_product():
    username = session.get('username')
    
    if not can_reserve(username):
        return jsonify({'error': f'❌ У пользователя {username} нет прав на снятие резерва'}), 403

    data = request.get_json()
    reserve_id = data.get('reserve_id')
    prod_id = data.get('product_id')

    if not reserve_id and not prod_id:
        return jsonify({'error': 'Не указан резерв или товар'}), 400

    conn = get_db()
    cur = conn.cursor()

    if reserve_id:
        cur.execute(
            "SELECT id, product_id, quantity, comment FROM reserves WHERE id = %s AND is_active = 1",
            (reserve_id,)
        )
        reserve = cur.fetchone()

        if not reserve:
            cur.close()
            conn.close()
            return jsonify({'error': 'Резерв не найден или уже снят'}), 400

        cur.execute("UPDATE reserves SET is_active = 0 WHERE id = %s", (reserve_id,))
        cur.execute(
            "INSERT INTO stock_moves (product_id, quantity, \"user\", comment) VALUES (%s, %s, %s, %s)",
            (reserve['product_id'], 0, username, 
             f'Снят резерв #{reserve_id} ({reserve["quantity"]} шт, пометка: {reserve["comment"]})')
        )
        conn.commit()
        cur.close()
        conn.close()
        return jsonify({
            'status': 'OK', 
            'message': f'Резерв #{reserve_id} снят ({reserve["quantity"]} шт)'
        })

    else:
        cur.execute(
            "SELECT id, quantity FROM reserves WHERE product_id = %s AND is_active = 1",
            (prod_id,)
        )
        reserves = cur.fetchall()

        if not reserves:
            cur.close()
            conn.close()
            return jsonify({'error': 'Нет активных резервов для этого товара'}), 400

        total = sum(r['quantity'] for r in reserves)
        count = len(reserves)

        cur.execute(
            "UPDATE reserves SET is_active = 0 WHERE product_id = %s AND is_active = 1",
            (prod_id,)
        )
        cur.execute(
            "INSERT INTO stock_moves (product_id, quantity, \"user\", comment) VALUES (%s, %s, %s, %s)",
            (prod_id, 0, username, f'Сняты все резервы ({count} шт): {total} ед.')
        )
        conn.commit()
        cur.close()
        conn.close()
        return jsonify({
            'status': 'OK', 
            'message': f'Снято {count} резервов на {total} единиц'
        })


@app.route('/api/reserves')
@login_required
def get_reserves():
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT 
            r.id,
            r.product_id,
            p.name as product_name,
            p.unit,
            r.quantity,
            r.user,
            r.comment,
            to_char(r.created_at + INTERVAL '5 hours', 'YYYY-MM-DD HH24:MI:SS') as created_at
        FROM reserves r
        JOIN products p ON r.product_id = p.id
        WHERE r.is_active = 1
        ORDER BY p.sort_order, p.id, r.created_at DESC
    ''')
    reserves = cur.fetchall()
    cur.close()
    conn.close()
    return jsonify([dict(row) for row in reserves])


# ============================================================
# ЭКСПОРТ В EXCEL
# ============================================================

@app.route('/export_excel')
@login_required
def export_excel():
    if session.get('role') != 'admin':
        return render_template('access_denied.html'), 403
    
    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        SELECT 
            p.id,
            p.name,
            COALESCE(SUM(sm.quantity), 0) as balance,
            COALESCE((
                SELECT SUM(r.quantity) 
                FROM reserves r 
                WHERE r.product_id = p.id AND r.is_active = 1
            ), 0) as reserved,
            p.unit
        FROM products p
        LEFT JOIN stock_moves sm ON p.id = sm.product_id
        GROUP BY p.id, p.name, p.unit, p.sort_order
        ORDER BY p.sort_order, p.id
    ''')
    data = cur.fetchall()
    cur.close()
    conn.close()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки склада"
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a3c5e", end_color="1a3c5e", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    data_font = Font(size=10)
    data_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    number_alignment = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    headers = ['№ п/п', 'Наименование позиции', 'Ед. изм.', 'Остаток, шт', 'Резерв, шт', 'Доступно, шт']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    for row_idx, item in enumerate(data, 2):
        available = item['balance'] - item['reserved']
        
        cell = ws.cell(row=row_idx, column=1, value=item['id'])
        cell.font = data_font
        cell.alignment = number_alignment
        cell.border = thin_border
        
        cell = ws.cell(row=row_idx, column=2, value=item['name'])
        cell.font = data_font
        cell.alignment = data_alignment
        cell.border = thin_border
        
        cell = ws.cell(row=row_idx, column=3, value=item['unit'] or 'шт')
        cell.font = data_font
        cell.alignment = number_alignment
        cell.border = thin_border
        
        cell = ws.cell(row=row_idx, column=4, value=item['balance'])
        cell.alignment = number_alignment
        cell.border = thin_border
        if item['balance'] < 5 and item['balance'] > 0:
            cell.font = Font(color="FF8F00", bold=True, size=10)
            cell.fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
        elif item['balance'] <= 0:
            cell.font = Font(color="FF0000", bold=True, size=10)
            cell.fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
        else:
            cell.font = Font(color="1B5E20", size=10)
        
        cell = ws.cell(row=row_idx, column=5, value=item['reserved'])
        cell.alignment = number_alignment
        cell.border = thin_border
        if item['reserved'] > 0:
            cell.font = Font(color="F57C00", bold=True, size=10)
            cell.fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
        else:
            cell.font = data_font
        
        cell = ws.cell(row=row_idx, column=6, value=available)
        cell.alignment = number_alignment
        cell.border = thin_border
        cell.font = Font(color="1B5E20", bold=True, size=10)
    
    column_widths = {'A': 8, 'B': 60, 'C': 12, 'D': 15, 'E': 15, 'F': 15}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    ws.freeze_panes = 'A2'
    
    total_items = len(data)
    low_items = len([item for item in data if 0 < item['balance'] < 5])
    zero_items = len([item for item in data if item['balance'] <= 0])
    reserved_items = len([item for item in data if item['reserved'] > 0])
    
    ws.append([])
    ws.append([f'Дата выгрузки: {get_perm_time().strftime("%d.%m.%Y %H:%M")} (Пермь)'])
    ws.append([f'Всего позиций: {total_items}'])
    ws.append([f'Позиций с остатком менее 5 шт: {low_items}'])
    ws.append([f'Позиций с нулевым остатком: {zero_items}'])
    ws.append([f'Позиций в резерве: {reserved_items}'])
    ws.append([f'Выгрузил: {session.get("username")} (администратор)'])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Остатки_склада_{get_perm_time().strftime('%Y%m%d_%H%M')}.xlsx"
    
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


# ============================================================
# ЗАПУСК
# ============================================================

if __name__ == '__main__':
    print("=" * 60)
    print("🚀 Запуск приложения...")
    print(f"📦 База данных: {DATABASE_URL[:50]}...")
    print(f"🔑 SECRET_KEY: {'*' * 20} (скрыт)")
    print("=" * 60)
    app.run(debug=True, host='0.0.0.0', port=5000)